#!/usr/bin/env python3
"""Freeze the displayed cell grid of MacLean et al. 2014 Tables S1–S7.

Run with a Python environment containing python-docx. The original Word files
remain authoritative; snapshots preserve their printed text and row order.
"""

import csv
from pathlib import Path

from docx import Document


FOLDER = Path(__file__).resolve().parent
EXPECTED = {1: (37, 6), 2: (24, 11), 3: (43, 3), 4: (568, 7),
            5: (37, 5), 6: (10, 8), 7: (46, 8)}


def main():
    for number, (expected_rows, expected_columns) in EXPECTED.items():
        source = FOLDER / f"pnas.1323533111.st0{number}.docx"
        document = Document(source)
        if len(document.tables) != 1:
            raise ValueError(f"{source.name}: expected one table")
        table = document.tables[0]
        if len(table.rows) != expected_rows:
            raise ValueError(f"{source.name}: unexpected row count")
        snapshot = FOLDER / f"MacLean_etal_2014_TableS{number}_snapshot.csv"
        with snapshot.open("w", newline="", encoding="utf-8") as stream:
            writer = csv.writer(stream)
            writer.writerow(["source_row"] + [f"c{i}" for i in range(1, expected_columns + 1)])
            for row_number, row in enumerate(table.rows, start=1):
                cells = [cell.text for cell in row.cells]
                if len(cells) != expected_columns:
                    raise ValueError(f"{source.name}: row {row_number} has {len(cells)} cells")
                writer.writerow([row_number] + cells)
        print(f"{snapshot.name}: {expected_rows} source rows")


if __name__ == "__main__":
    main()
