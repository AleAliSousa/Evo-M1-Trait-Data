"""Myhrvold_etal_2015_AmnioteDatabaseAug2015.py -- reformat the amniote life-history
database (Myhrvold et al. 2015) into a long analysis table with per-value references.

Registry note: the registry's "Item name" for this item is "Myhrvold_etal_2015_AmnioteDatabaseAug2015"
(no file extension, matching house convention). An earlier version of this row had the
original Data_Files filename (with ".csv") leaking into the registry's "Item number"
column (plus, once, stray XML markup from a bad paste) -- both were fixed directly in
__ReadMe.xlsx; this script's own output CSV name is independently ITEM_NAME + ".csv".

Frozen sources (byte-identical copies of the two Data_Files, never edited):
  Myhrvold_etal_2015_AmnioteDatabaseAug2015_snapshot.csv            -- values, -999 = missing
  Myhrvold_etal_2015_AmnioteDatabaseReferencesAug2015_snapshot.csv  -- same shape/row order,
                                                                       per-cell reference string

Run: python3 Myhrvold_etal_2015_AmnioteDatabaseAug2015.py
"""
import os
import pandas as pd

HERE = os.path.dirname(os.path.abspath(__file__))
ITEM_NAME = "Myhrvold_etal_2015_AmnioteDatabaseAug2015"

# ---- repo root: walk up from HERE to find __ReadMe.xlsx -----------------------------
def find_repo_root(start):
    d = start
    while os.path.dirname(d) != d and not os.path.exists(os.path.join(d, "__ReadMe.xlsx")):
        d = os.path.dirname(d)
    return d if os.path.exists(os.path.join(d, "__ReadMe.xlsx")) else None

REPO = find_repo_root(HERE)

TAX_COLS = ["class", "order", "family", "genus", "species", "subspecies", "common_name"]

db = pd.read_csv(os.path.join(HERE, "Myhrvold_etal_2015_AmnioteDatabaseAug2015_snapshot.csv"))
refs = pd.read_csv(os.path.join(HERE, "Myhrvold_etal_2015_AmnioteDatabaseReferencesAug2015_snapshot.csv"),
                    low_memory=False)
assert list(db.columns) == list(refs.columns)
assert (db[TAX_COLS].astype(str).values == refs[TAX_COLS].astype(str).values).all(), \
    "Database and References files must be row-for-row, taxon-for-taxon aligned."

TRAIT_COLS = [c for c in db.columns if c not in TAX_COLS]

field_info = pd.read_csv(
    os.path.join(REPO, "Myhrvold_etal_2015", "Supplemental_Materials",
                  "Supplemental_Table_9_Dataset_Field_Information.csv"),
    skiprows=1,
)
field_info.columns = ["field_number", "field_name", "units", "taxon_group", "description"]
field_info = field_info.dropna(subset=["field_name"])
unit_map = dict(zip(field_info["field_name"], field_info["units"]))

db_long = db.melt(id_vars=TAX_COLS, value_vars=TRAIT_COLS, var_name="trait", value_name="value")
refs_long = refs.melt(id_vars=TAX_COLS, value_vars=TRAIT_COLS, var_name="trait", value_name="reference")
assert len(db_long) == len(refs_long)

long = db_long.copy()
long["reference"] = refs_long["reference"]

## -999 is the paper's own missing-value sentinel (see Supplemental_Table_9 note /
## metadata.htm); drop rows with no value. Subspecies -999 means "no subspecies
## recorded" (not a missing value) and is kept as such, turned to NA for clarity.
long["value"] = pd.to_numeric(long["value"], errors="coerce")
long = long[(long["value"].notna()) & (long["value"] != -999)].copy()
long["reference"] = long["reference"].replace("-999", pd.NA)
long["subspecies"] = long["subspecies"].replace(-999, pd.NA).replace("-999", pd.NA)

long["unit"] = long["trait"].map(unit_map)
## data_role = secondary throughout: Myhrvold et al. 2015 is itself a compilation --
## every value's reference field names the paper(s) it was actually measured/reported
## in (or a derivation, e.g. "mean of X and Y", "X from median of N(...)").
long["data_role"] = "secondary"
long["source_location"] = ("Amniote_Database_Aug_2015.csv + "
                            "Amniote_Database_References_Aug_2015.csv (per-cell reference)")

out = long[TAX_COLS + ["trait", "unit", "value", "reference", "data_role", "source_location"]]
out = out.sort_values(["class", "order", "family", "genus", "species", "trait"]).reset_index(drop=True)

out_path = os.path.join(HERE, f"{ITEM_NAME}.csv")
out.to_csv(out_path, index=False)
print(f"{ITEM_NAME}.csv: {len(out)} rows written")

## ---- public TSV ----------------------------------------------------------------
if REPO:
    import openpyxl
    wb = openpyxl.load_workbook(os.path.join(REPO, "__ReadMe.xlsx"), data_only=True)
    ws = wb["Sheet1"]
    hdr = [c.value for c in ws[1]]
    name_col = hdr.index("Item name") + 1
    enc_col = hdr.index("Item encoded") + 1
    item_encoded = None
    for r in range(2, ws.max_row + 1):
        if ws.cell(row=r, column=name_col).value == ITEM_NAME:
            item_encoded = ws.cell(row=r, column=enc_col).value
            break
    tsv_dir = os.path.join(REPO, "__Public", "comparative-data")
    if not item_encoded:
        print(f"WARNING: no 'Item encoded' for '{ITEM_NAME}' in __ReadMe.xlsx; TSV skipped.")
    elif not os.path.isdir(tsv_dir):
        print(f"WARNING: shared folder not found: {tsv_dir}; TSV skipped.")
    else:
        tsv_path = os.path.join(tsv_dir, f"{item_encoded}.tsv")
        out.to_csv(tsv_path, sep="\t", index=False)
        print(f"Wrote TSV: {item_encoded}")
else:
    print("WARNING: repo root (__ReadMe.xlsx) not found; TSV skipped.")
