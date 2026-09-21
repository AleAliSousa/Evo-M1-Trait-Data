#!/usr/bin/env python3
"""Build one journal-faithful XLSX snapshot from a Zilles et al. 2011 XLS table.
Usage: python3 Zilles_etal_2011_build_snapshot.py Zilles_etal_2011_TableS1
"""
from pathlib import Path
from copy import copy
import sys, xlrd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
CFG={
 'Zilles_etal_2011_TableS1':('nyas_5978_sm_table1.xls',0,1,9,'TableS1'),
 'Zilles_etal_2011_TableS2':('nyas_5978_sm_table2.xls',2,5,11,'TableS2'),
 'Zilles_etal_2011_TableS3':('nyas_5978_sm_table3.xls',2,5,10,'TableS3'),
}
def main(stem):
 if stem not in CFG: raise SystemExit('Unknown item: '+stem)
 src_name,title_row,header_row,ncol,sheet=CFG[stem]
 folder=Path(__file__).resolve().parent; src=folder/src_name
 if not src.exists(): raise SystemExit('Missing source: '+str(src))
 sh=xlrd.open_workbook(src).sheet_by_index(0)
 title=str(sh.cell_value(title_row,0)).strip()
 headers=[str(sh.cell_value(header_row,c)).replace('\n',' ').strip() for c in range(ncol)]
 rows=[]
 for r in range(header_row+1,sh.nrows):
  vals=[]
  for c in range(ncol):
   cell=sh.cell(r,c); v=cell.value
   if cell.ctype==xlrd.XL_CELL_NUMBER and float(v).is_integer(): v=int(v)
   if isinstance(v,str): v=v.replace('\xa0',' ').strip()
   vals.append(v if v!='' else None)
  if any(v is not None for v in vals): rows.append(vals)
 wb=Workbook(); ws=wb.active; ws.title=sheet; ws.sheet_view.showGridLines=False
 ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=ncol); ws.cell(1,1,title)
 ws.cell(1,1).font=Font(bold=True); ws.cell(1,1).alignment=Alignment(wrap_text=True)
 for c,h in enumerate(headers,1):
  x=ws.cell(2,c,h); x.font=Font(bold=True); x.fill=PatternFill('solid',fgColor='D9EAF7'); x.alignment=Alignment(wrap_text=True,vertical='bottom'); x.border=Border(top=Side(style='thin'),bottom=Side(style='thin'))
 for r,row in enumerate(rows,3):
  for c,v in enumerate(row,1): ws.cell(r,c,v)
 ws.freeze_panes='A3'; ws.auto_filter.ref=f'A2:{get_column_letter(ncol)}{len(rows)+2}'
 ws.column_dimensions['A'].width=28
 for c in range(2,ncol+1): ws.column_dimensions[get_column_letter(c)].width=20
 ws.row_dimensions[1].height=30; ws.row_dimensions[2].height=55
 ws.page_setup.orientation='landscape'; ws.sheet_properties.pageSetUpPr.fitToPage=True; ws.page_setup.fitToWidth=1
 out=folder/f'{stem}_snapshot.xlsx'; wb.save(out); print(f'Wrote {out.name}: {len(rows)} nonblank source rows')
if __name__=='__main__':
 if len(sys.argv)!=2: raise SystemExit('Supply one item stem')
 main(sys.argv[1])
