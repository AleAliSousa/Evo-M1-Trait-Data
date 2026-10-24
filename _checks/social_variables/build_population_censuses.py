"""Rebuild source tables from preserved author archives and verified primary passages.
Does not change source snapshots, pool distinct studies, or infer counts from labels.
"""
import csv, json, hashlib, collections, statistics, io, zipfile
from pathlib import Path
import xml.etree.ElementTree as ET
import openpyxl
ROOT=Path(__file__).resolve().parents[2]
def write(path,rows):
 path=ROOT/path;path.parent.mkdir(parents=True,exist_ok=True)
 assert rows
 with path.open('w',newline='') as f:
  w=csv.DictWriter(f,list(rows[0]));w.writeheader();w.writerows(rows)
def definitions(folder,code,text,role):
 write(Path(folder)/'reference_tables'/f'{folder}_definitions.csv',[dict(Code=code,Definition=text,role=role)])
# Retain population rows and original reference strings; no taxon substitutions from phylogeny.
f='Jaeggi_etal_2020';w=openpyxl.load_workbook(io.BytesIO(zipfile.ZipFile(ROOT/f/'source/author_workbook_original.zip').read('Artiodactyla_Dataset_corrected_11Mar.xlsx')),data_only=True)
s=w['Population data for analysis'];headers=[str(c.value) for c in s[7][:30]]
raw=[]
for row in s.iter_rows(min_row=8,max_col=30):
 if row[0].value:
  raw.append(dict(source_row=row[0].row,**{k:c.value for k,c in zip(headers,row)}))
assert len(raw)==245
write(Path(f)/f'{f}_Population_snapshot.csv',raw)
pops=[]
for r in raw:
 pops.append(dict(Species=r['Genus_species'].replace('_',' ').strip(),population_id=f'Jaeggi2020_row{r["source_row"]}',population=r['Population'],subspecies_printed=r['Subspecies'],social_state=str(r['Social state - for R']).strip(),detailed_state=str(r['VASO - for R']).strip(),supporting_references=r['References'],role='secondary',primary_paper_verified=False,age_scope='adult male-female social organization',quality_flag='Compiler classification from cited field papers; not a numerical group count'))
write(Path(f)/f'{f}_Populations.csv',pops)
def species_summary(folder,populations):
 groups=collections.defaultdict(list)
 for r in populations:
  if not r.get('merge_eligible',True):continue
  groups[r['Species']].append(r)
 out=[]
 for sp,rs in sorted(groups.items()):
  out.append(dict(Species=sp,social_states=' | '.join(sorted({r['social_state'] for r in rs})),n_population_records=len(rs),population_ids='; '.join(r['population_id'] for r in rs),supporting_references='; '.join(sorted({str(r['supporting_references']) for r in rs})),role='secondary',statistic='set of reported population states; not a modal species classification'))
 write(Path(folder)/f'{folder}_SocialStates.csv',out)
 return out
ja=species_summary(f,pops);assert len(ja)==100
definitions(f,'social_states','Set of social-organization states reported across the source populations. Adult male-female composition, not mating system and not numerical group size. All 245 population rows and references retained. Secondary synthesis of primary field studies; supporting papers have not all been independently checked. States are unioned, never averaged or reduced to the first/modal state. Species names come from the source table, not phylogenetic proxy substitutions.','secondary')
# Public GitHub S5 archive is a recoverable SUBSET, not the inaccessible full Dryad workbook.
f='Qiu_etal_2022';z=zipfile.ZipFile(ROOT/f/'source/Model_S5_author_archive.zip');w=openpyxl.load_workbook(io.BytesIO(z.read(next(n for n in z.namelist() if n.endswith('/S5.xlsx')))),data_only=True)
raw=[]
for s in w:
 for i,r in enumerate(s.values,1):
  if r[0] and '_' in str(r[0]) and r[0]!='Genus_species':
   raw.append(dict(source_sheet=s.title,source_row=i,Species_printed=r[0],habitat_heterogeneity=r[1],social_state_printed=r[2],n_studies=r[3],body_mass_g=r[4],supporting_references=r[5]))
assert len(raw)==141 and len({r['Species_printed'] for r in raw})==65
write(Path(f)/f'{f}_S5_snapshot.csv',raw)
codes={'S':'Solitary','P':'Pair-living','G':'Stable group','VG':'Variable group','SP':'Solitary + pair','PG':'Pair + group','SPG':'Solitary + pair + group','SG':'Solitary + group'}
pops=[]
for r in raw:
 code=str(r['social_state_printed']).strip();assert code in codes
 pops.append(dict(Species=r['Species_printed'].replace('_',' ').strip(),population_id=f'Qiu2022_{r["source_sheet"]}_row{r["source_row"]}',population='Not provided in S5 analysis workbook; recover from full Dryad table',social_state=codes[code],supporting_references=r['supporting_references'],source_sheet=r['source_sheet'],role='secondary',primary_paper_verified=False,age_scope='adults only; juveniles excluded',merge_eligible=code!='SG',quality_flag='Incomplete archive: 141 records versus 149 reported; preserve excluded-from-model rows; not all unique populations reverified'+('; SG contradicts published statement that SG did not occur: held from merge' if code=='SG' else '')))
write(Path(f)/f'{f}_Populations.csv',pops)
qi=species_summary(f,pops)
definitions(f,'social_states','Set of adult social-organization states in the recoverable author S5 analysis archive (141 records, 65 printed species), including records excluded from the model for missing covariates. Incomplete relative to the reported 149 populations. Exact population locations are absent from this archive. One SG record conflicts with the publication and is held from the merge. Secondary field-literature synthesis, not new primary observations; citations retained. Counts of records are not numbers of individuals. States unioned without choosing a modal state.','secondary')
# Rabbit site means re-reported by the originating study team; 2015 source explicitly retained.
f='Ziege_etal_2016';tree=ET.parse(ROOT/f/'source/PMC4908761.xml');tab=next(x for x in tree.iter('table-wrap') if x.attrib.get('id')=='Tab1')
rows=[];setting=''
for tr in tab.findall('.//tbody/tr'):
 cells=[''.join(x.itertext()).strip() for x in tr]
 if len(cells)==1:setting=cells[0];continue
 if len(cells)!=7:continue
 rows.append(dict(Species='Oryctolagus cuniculus',population=cells[0],coordinates='; '.join(cells[1:3]),site_area_ha=cells[3],urbanity=cells[4],population_density_per_ha=cells[5],group_size_mean=float(cells[-1]),setting=setting,source_location='Table 1',original_source_doi='10.1111/jzo.12207',reporting_source_doi='10.1186/s12898-016-0083-y',role='primary',provenance='Same originating team re-reporting their field census; not an independent 2016 replicate',age_scope='Unspecified in reprinted table; do not assume adult-only',method='Behavioral observations and ferrets driving members from burrows',quality_flag='Underlying 2015 paper not independently retrieved; population summaries retained as reported'))
# Table cells may have split latitude/longitude columns; fail rather than misalign.
if len(rows)!=15:
 print('Rabbit table row layout:',[[''.join(x.itertext()).strip() for x in tr] for tr in tab.findall('.//tbody/tr')][:5]);raise AssertionError(len(rows))
write(Path(f)/f'{f}_Population_snapshot.csv',rows)
summary=dict(Species='Oryctolagus cuniculus',mean_of_population_means=statistics.mean(r['group_size_mean'] for r in rows),n_populations=15,population_mean_min=min(r['group_size_mean'] for r in rows),population_mean_max=max(r['group_size_mean'] for r in rows),statistic='unweighted mean of 15 reported population means; NOT group-weighted mean',role='primary',original_source_doi='10.1111/jzo.12207',reporting_source_doi='10.1186/s12898-016-0083-y')
write(Path(f)/f'{f}_Summary.csv',[summary])
definitions(f,'mean_of_population_means','Unweighted arithmetic mean of 15 reported mean social-group sizes around Frankfurt, Germany (rural, suburban and urban). Individual site means and coordinates are retained in the snapshot; underlying group sample sizes are unavailable, so this is not a group-weighted species mean. Re-reported by the originating study team in Table 1 (2016), explicitly citing Ziege et al. 2015 DOI 10.1111/jzo.12207; count the team only once. Behavioral observations plus burrow censuses using ferrets; age restriction unspecified.','primary')
# Direct author colony estimate, not a precise census or mean social unit.
f='Modlinska_Stryjek_2016';tree=ET.parse(ROOT/f/'source/PMC4890768.xml')
assert any('50 individuals' in ''.join(p.itertext()) for p in tree.iter('p'))
row=dict(Species='Rattus norvegicus',colony_size_estimate=50,approximate=True,population='Farm on outskirts of Warsaw, Poland',statistic='approximate whole-colony population estimate',method='Six months camera traps plus burrows, paths, droppings and bait consumption; also prior same-team study',age_scope='not specified',source_location='Materials and Methods: Animals',source_doi='10.1371/journal.pone.0156741',role='primary',quality_flag='Population colony may contain multiple social subgroups; not a census-derived social-group mean')
write(Path(f)/f'{f}_Colony_snapshot.csv',[row]);write(Path(f)/f'{f}_Colony.csv',[row])
definitions(f,'colony_size_estimate','Approximately 50 rats in one free-living farm colony outside Warsaw, estimated by the authors using six months of camera traps and field signs, also drawing on a prior same-team study. Whole-colony population estimate, not a precise head count or social-subgroup mean. Age composition unspecified. Keep approximation and commensal farm context.','primary')
# Ferret den occupancy, no invented individual den records.
f='Norbury_etal_1998'
raw=[dict(Species='Mustela furo',occupants_per_den=1,number_of_excavated_dens=39,derivation='41 examined minus two shared',source_location='p.154 Den sharing by ferrets'),dict(Species='Mustela furo',occupants_per_den=2,number_of_excavated_dens=2,derivation='Two dens each contained collared ferret and another adult male',source_location='p.154 Den sharing by ferrets')]
write(Path(f)/f'{f}_DenCensus_snapshot.csv',raw)
row=dict(Species='Mustela furo',mean_den_occupancy=43/41,n_examined_dens=41,n_shared_dens=2,mean_formula='(39*1 + 2*2)/41',shared_fraction=2/41,setting='Feral ferrets; South Island New Zealand pastoral short-tussock habitat',statistic='derived mean occupants per excavated occupied den',role='primary',age_scope='Collared animals subadult/adult; shared pairs adult males',quality_flag='Dens selected because collared ferrets present; repeated animals possible; not independent social groups. Later telemetry sharing excluded from this denominator.')
write(Path(f)/f'{f}_DenSummary.csv',[row])
definitions(f,'mean_den_occupancy','Derived mean 43/41 occupants per excavated occupied ferret den: 39 single occupants and two pairs in 41 excavations (Norbury et al. 1998 p.154). Selected dens contained radio-collared ferrets; not a random sample or necessarily independent animals. Feral Mustela furo, mapped to the project ferret anchor Mustela putorius furo. Daytime den occupancy is distinct from stable social-group size; later telemetry sharing is outside this denominator.','primary')
for f in ['Jaeggi_etal_2020','Qiu_etal_2022','Ziege_etal_2016','Modlinska_Stryjek_2016','Norbury_etal_1998']:
 manifests=[]
 for src in sorted((ROOT/f/'source').iterdir()):
  # Office display copies may gain local metadata; immutable archives are authoritative.
  if src.suffix in ('.xlsx','.docx'):continue
  manifests.append(dict(file=src.name,bytes=src.stat().st_size,sha256=hashlib.sha256(src.read_bytes()).hexdigest(),retrieved='2026-10-23'))
 write(Path(f)/'source_manifest.csv',manifests)
print('Built',len(ja),'artiodactyl species;',len(qi),'marsupial species; rabbit, rat and ferret summaries')
