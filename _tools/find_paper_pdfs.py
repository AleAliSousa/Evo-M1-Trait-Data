#!/usr/bin/env python3
"""
Find and fetch the paper PDF for every paper folder that has none.

Pipeline (each stage only runs for folders still unresolved):

  1. audit      read _checks/folder_audit.csv, keep paper_pdf == 0
  2. identify   folder -> full citation / DOI, from (in priority order)
                  a. _checks/paper_citations.csv         manual overrides
                  b. __ReadMe.xlsx registry              (Item encoded prefix, DOI, title)
                  c. Sensory reconciliation + references (restricted repo)
  3. endnote    match in the EndNote library via _tools/endnote.py
                  DOI -> exact;  first-author + year -> title-overlap score.
                  Only 'strong' matches (>= 0.5 title overlap, or DOI) are copied.
  4. crossref   resolve a DOI from the citation string (api.crossref.org)
  5. oa         open-access PDF: Europe PMC (PMCID -> render), Unpaywall
                  (only if UNPAYWALL_EMAIL is set), Semantic Scholar openAccessPdf
  6. local      scan --scan-dir folders for PDFs; filename author+year and
                  first-page title overlap (needs pypdfium2)
  7. report     _checks/paper_pdf_sources.csv (+ .md) — one row per folder:
                  status, source kind, exact source path/URL, EndNote id, DOI,
                  confidence, manual links (doi.org / ResearchGate / Scholar)

Nothing is copied unless --apply is given. Copies are named <Folder>.pdf and
recorded in _checks/paper_pdf_provenance.csv (append-only).

Usage
    find_paper_pdfs.py                       # dry run, report only
    find_paper_pdfs.py --apply               # copy / download into folders
    find_paper_pdfs.py --apply --scan-dir ~/Documents --scan-dir ~/Desktop
    find_paper_pdfs.py --folder Heffner_etal_2001 --apply
    find_paper_pdfs.py --no-online           # EndNote + local only

Environment
    ENDNOTE_DATA      override References.Data location (see endnote.py)
    UNPAYWALL_EMAIL   enable the Unpaywall lookup (they require a contact email)
    EVOM1_RESTRICTED  path to Evo-M1-Trait-Data-restricted (default: sibling
                      OneDrive folder)

House rules encoded here (learned 2026-09-26):
  * folder_audit joins registry rows by *Item name prefix*, so Foo_2001_b
    inherits Foo_2001's DOI/title. Suffix folders (_a/_b/_c) are therefore
    never trusted to the registry DOI; the sensory reconciliation sheet or a
    manual override in paper_citations.csv decides.
  * The audit's supplement regex matches 'tables' (TableS, case-insensitive);
    a paper PDF named '...book tables.pdf' is misclassified as supp_pdf.
  * Author+year alone is NOT a match for Heffner/Koay-lab papers (several per
    year); title overlap must confirm, and 'poor' matches are reported, not
    copied.
  * OneDrive Files-On-Demand placeholders raise 'Operation timed out' on read;
    the scanner skips them and lists them as unchecked.
"""

from __future__ import annotations

import argparse
import csv
import glob
import hashlib
import json
import os
import re
import shutil
import sys
import time
import urllib.parse
import urllib.request
import concurrent.futures as cf

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
CHECKS = os.path.join(ROOT, "_checks")
sys.path.insert(0, HERE)
import endnote as en  # noqa: E402  (read-only EndNote access)

RESTRICTED = os.environ.get(
    "EVOM1_RESTRICTED",
    os.path.join(os.path.dirname(os.path.dirname(ROOT)), "Evo-M1-Trait-Data-restricted"),
)
SENSORY = os.path.join(RESTRICTED, "unpublished_data",
                       "____Unpublished__Sensory_audiovisual", "SensoryData_compiled_check")

STOP = set("the of in and a an to for from on by with at its is are as i ii iii iv hearing".split())
UA = {"User-Agent": "Evo-M1-Trait-Data find_paper_pdfs (research data curation)"}


# ---------------------------------------------------------------- helpers

def toks(s: str) -> set[str]:
    return {w for w in re.findall(r"[a-z]{4,}", (s or "").lower()) if w not in STOP}


def parse_folder(f: str):
    """Foo_Bar_1999_b -> ('Foo', 'Bar'|'etal'|'', '1999', 'b'|None)."""
    m = re.match(r"^([A-Za-z\-]+)_([A-Za-z\-]*)_(\d{4})(?:_([a-z]))?$", f)
    return (m.group(1), m.group(2), m.group(3), m.group(4)) if m else (None, None, None, None)


def doi_in(s: str) -> str:
    m = re.search(r"10\.\d{4,9}/[^\s\"”,;]+", s or "")
    return m.group().rstrip(".") .replace(".supp", "") if m else ""


def read_with_timeout(fn, timeout, default):
    with cf.ThreadPoolExecutor(1) as ex:
        fut = ex.submit(fn)
        try:
            return fut.result(timeout=timeout)
        except Exception:
            return default


def http_json(url, params=None, timeout=30):
    if params:
        url += ("&" if "?" in url else "?") + urllib.parse.urlencode(params)
    req = urllib.request.Request(url, headers=UA)
    with urllib.request.urlopen(req, timeout=timeout) as r:
        return json.loads(r.read().decode("utf-8", "replace"))


def http_bytes(url, timeout=90):
    req = urllib.request.Request(url, headers=UA)
    with urllib.request.urlopen(req, timeout=timeout) as r:
        return r.read(), r.headers.get("Content-Type", "")


def is_pdf(b: bytes) -> bool:
    return b[:5] == b"%PDF-"


# ---------------------------------------------------------------- stage 1: audit

def load_audit(only=None):
    p = os.path.join(CHECKS, "folder_audit.csv")
    rows = list(csv.DictReader(open(p, newline="", encoding="utf-8")))
    out = [r for r in rows if r["paper_pdf"] == "0" and r["folder"] != "reconstructed_results"
           and parse_folder(r["folder"])[0]]
    if only:
        out = [r for r in out if r["folder"] in only]
    return out


# ---------------------------------------------------------------- stage 2: identify

def load_overrides():
    p = os.path.join(CHECKS, "paper_citations.csv")
    if not os.path.exists(p):
        return {}
    return {r["folder"]: r for r in csv.DictReader(open(p, newline="", encoding="utf-8"))}


def load_registry():
    try:
        import openpyxl  # noqa
        import pandas as pd
    except ImportError:
        return {}
    xl = glob.glob(os.path.join(ROOT, "__ReadMe.xlsx"))
    if not xl:
        return {}
    s1 = pd.read_excel(xl[0], sheet_name="Sheet1")
    s1.columns = [c.strip() for c in s1.columns]
    reg = {}
    for _, r in s1.iterrows():
        item = str(r.get("Item encoded", ""))
        key = re.sub(r"_(Table|Fig|Text|Appendix|Supp|Data|Suppl|Electronic).*$", "", item)
        if not parse_folder(key)[0]:
            continue
        reg.setdefault(key, dict(
            cite=str(r.get("Citation (APA 7th-Annotated)", "") or ""),
            doi=doi_in(str(r.get("DOI (or Alt)", "") or "")),
            title=str(r.get("Item full original title", "") or ""),
        ))
    return reg


def load_sensory():
    """folder -> full citation via reconciliation + references sheets."""
    rec_p = os.path.join(SENSORY, "Sensory_reference_reconciliation.csv")
    ref_p = os.path.join(SENSORY, "SensoryData_compiled_references.csv")
    if not (os.path.exists(rec_p) and os.path.exists(ref_p)):
        return {}
    refs = read_with_timeout(lambda: list(csv.DictReader(open(ref_p, newline="", encoding="utf-8-sig"))), 20, None)
    rec = read_with_timeout(lambda: list(csv.DictReader(open(rec_p, newline="", encoding="utf-8-sig"))), 20, None)
    if refs is None or rec is None:
        print("  ! sensory sheets are un-hydrated cloud placeholders; skipped", file=sys.stderr)
        return {}
    refmap = {r["reference_short"].strip(): r["full_citation"] for r in refs}
    out = {}
    for r in rec:
        folders = [x.strip() for x in (r.get("Paper folder") or "").split(";")]
        codes = [x.strip() for x in (r.get("Matched Reference-sheet code") or "").split(";")]
        if len(codes) == 1 and len(folders) > 1:
            codes = codes * len(folders)
        for f, c in zip(folders, codes):
            if f and refmap.get(c):
                out.setdefault(f, dict(code=c, cite=refmap[c], issue=r.get("Issue", "")))
    return out


def identify(folder, overrides, registry, sensory):
    """Return dict(cite, doi, title, id_source)."""
    a1, a2, yr, suf = parse_folder(folder)
    if folder in overrides:
        o = overrides[folder]
        return dict(cite=o.get("citation", ""), doi=doi_in(o.get("doi", "")) or doi_in(o.get("citation", "")),
                    title=o.get("title", ""), id_source="paper_citations.csv")
    if folder in sensory:
        s = sensory[folder]
        return dict(cite=s["cite"], doi=doi_in(s["cite"]), title="", id_source="sensory_reconciliation")
    if folder in registry and not suf:      # suffix folders inherit wrongly via prefix-join
        r = registry[folder]
        return dict(cite=r["cite"], doi=r["doi"], title=r["title"], id_source="registry")
    return dict(cite="", doi="", title="", id_source="none")


# ---------------------------------------------------------------- stage 3: endnote

def endnote_match(folder, ident, con):
    a1, a2, yr, suf = parse_folder(folder)
    desc = toks(" ".join([ident["cite"], ident["title"]]))
    rows, how = [], ""
    if ident["doi"]:
        rows = con.execute("SELECT * FROM refs WHERE trash_state=0 AND electronic_resource_number LIKE ?",
                           [f"%{ident['doi']}%"]).fetchall()
        how = "doi"
    if not rows:
        rows = con.execute("SELECT * FROM refs WHERE trash_state=0 AND author LIKE ? AND year LIKE ?",
                           [f"{a1}%", f"%{yr}%"]).fetchall()
        how = "author+year"
    if not rows:
        return None
    scored = []
    for r in rows:
        ct = toks(r["title"])
        ov = len(ct & desc) / max(1, len(ct)) if desc else 0.0
        if a2 and a2 != "etal" and a2.lower() not in (r["author"] or "").lower():
            ov *= 0.5                                   # co-author named in folder but absent
        scored.append((ov, r))
    scored.sort(key=lambda x: -x[0])
    ov, r = scored[0]
    pdfs = en.attachments(con, r["id"])
    if how == "doi":
        conf = "doi"
    elif ov >= 0.5:
        conf = "strong"
    elif ov >= 0.25:
        conf = "weak"
    else:
        conf = "poor"
    return dict(en_id=r["id"], en_cite=f"{en.authors(r['author'], 3)} ({r['year']})", en_title=r["title"],
                en_doi=r["electronic_resource_number"] or "", pdf=pdfs[0] if pdfs else "",
                overlap=round(ov, 2), confidence=conf, how=how,
                alternates="; ".join(f"[{x[1]['id']}] {x[1]['title'][:60]}" for x in scored[1:4]))


# ---------------------------------------------------------------- stage 4/5: online

def crossref_doi(cite, a1, yr):
    if not cite:
        return ""
    for attempt in range(3):
        try:
            j = http_json("https://api.crossref.org/works",
                          {"query.bibliographic": cite[:400], "rows": 4})
            for it in j["message"]["items"]:
                fam = (it.get("author") or [{}])[0].get("family", "")
                y = (it.get("issued", {}).get("date-parts") or [[None]])[0][0]
                if a1.lower()[:4] in fam.lower() and y and abs(int(y) - int(yr)) <= 1:
                    return it["DOI"]
            return ""
        except Exception as e:                      # 429 etc.
            if "429" in str(e):
                time.sleep(3 * (attempt + 1))
                continue
            return ""
    return ""


def oa_pdf(doi):
    """Return (url, source) of an open-access PDF or ('', reason)."""
    reasons = []
    # Europe PMC: PMCID + full text flags, then render endpoint
    try:
        j = http_json("https://www.ebi.ac.uk/europepmc/webservices/rest/search",
                      {"query": f"DOI:{doi}", "format": "json", "resultType": "lite"})
        hits = j.get("resultList", {}).get("result", [])
        if hits:
            h = hits[0]
            pmcid = h.get("pmcid")
            if pmcid and (h.get("isOpenAccess") == "Y" or h.get("inEPMC") == "Y"):
                return f"https://europepmc.org/articles/{pmcid}?pdf=render", f"europepmc:{pmcid}"
            reasons.append(f"europepmc: pmcid={pmcid} oa={h.get('isOpenAccess')}")
        else:
            reasons.append("europepmc: no record")
    except Exception as e:
        reasons.append(f"europepmc: {type(e).__name__}")
    # Unpaywall (contact email required by their terms)
    email = os.environ.get("UNPAYWALL_EMAIL")
    if email:
        try:
            j = http_json(f"https://api.unpaywall.org/v2/{doi}", {"email": email})
            loc = j.get("best_oa_location") or {}
            u = loc.get("url_for_pdf") or ""
            if u:
                return u, f"unpaywall:{loc.get('host_type', '')}"
            reasons.append(f"unpaywall: is_oa={j.get('is_oa')}")
        except Exception as e:
            reasons.append(f"unpaywall: {type(e).__name__}")
    else:
        reasons.append("unpaywall: skipped (UNPAYWALL_EMAIL unset)")
    # Semantic Scholar
    try:
        j = http_json(f"https://api.semanticscholar.org/graph/v1/paper/DOI:{doi}",
                      {"fields": "openAccessPdf"})
        u = (j.get("openAccessPdf") or {}).get("url") or ""
        if u:
            return u, "semantic_scholar"
        reasons.append("s2: no openAccessPdf")
    except Exception as e:
        reasons.append(f"s2: {type(e).__name__}")
    return "", "; ".join(reasons)


# ---------------------------------------------------------------- stage 6: local scan

def scan_local(dirs, targets, exclude_substrings):
    """targets: list of (folder, a1, yr, title_tokens). Returns hits per folder."""
    try:
        import pypdfium2 as pdfium
    except ImportError:
        print("  ! pypdfium2 not installed; local scan uses filenames only", file=sys.stderr)
        pdfium = None
    pdfs = []
    for d in dirs:
        for dp, dn, fn in os.walk(os.path.expanduser(d)):
            dn[:] = [x for x in dn if not x.startswith(".") and x not in ("node_modules", "Library")]
            for f in fn:
                p = os.path.join(dp, f)
                if f.lower().endswith(".pdf") and not any(s in p for s in exclude_substrings):
                    pdfs.append(p)
    print(f"  local scan: {len(pdfs)} PDFs under {len(dirs)} dir(s)", file=sys.stderr)

    def first_text(p):
        if pdfium is None:
            return ""
        def _r():
            d = pdfium.PdfDocument(p)
            return "\n".join(d[i].get_textpage().get_text_range() for i in range(min(2, len(d))))
        return read_with_timeout(_r, 15, None)

    unchecked, hits = [], {}
    for p in pdfs:
        tx = first_text(p)
        if tx is None:
            unchecked.append(p)
            tx = ""
        fn = os.path.basename(p).lower()
        low = tx.lower()[:4000]
        for folder, a1, yr, tt in targets:
            fn_hit = a1.lower() in fn and yr in fn
            auth = a1.lower() in low and yr in low
            ov = len(tt & toks(low)) / max(1, len(tt)) if tt else 0
            # require author evidence AND title evidence: generic short titles collide
            if (fn_hit and ov >= 0.5) or (auth and ov >= 0.75 and len(tt) >= 3):
                # filename author+year is the stronger evidence: rank it first
                hits.setdefault(folder, []).append((round(ov + (0.5 if fn_hit else 0), 2), p))
    return hits, unchecked


# ---------------------------------------------------------------- copy / provenance

def record_provenance(row):
    p = os.path.join(CHECKS, "paper_pdf_provenance.csv")
    new = not os.path.exists(p)
    with open(p, "a", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=["date", "folder", "dest", "source_kind", "source", "endnote_id",
                                           "doi", "confidence", "size_bytes", "md5"])
        if new:
            w.writeheader()
        w.writerow(row)


def place_pdf(folder, data_or_path, source_kind, source, apply, endnote_id="", doi="", confidence=""):
    dst_dir = os.path.join(ROOT, folder)
    dst = os.path.join(dst_dir, f"{folder}.pdf")
    if os.path.exists(dst):
        return "already-present", dst
    if not apply:
        return "would-copy", dst
    os.makedirs(dst_dir, exist_ok=True)
    if isinstance(data_or_path, bytes):
        open(dst, "wb").write(data_or_path)
    else:
        shutil.copy2(data_or_path, dst)
    b = open(dst, "rb").read()
    record_provenance(dict(date=time.strftime("%Y-%m-%d"), folder=folder, dest=os.path.basename(dst),
                           source_kind=source_kind, source=source, endnote_id=endnote_id, doi=doi,
                           confidence=confidence, size_bytes=len(b), md5=hashlib.md5(b).hexdigest()))
    return "copied", dst


# ---------------------------------------------------------------- main

def manual_links(folder, ident, doi):
    a1, a2, yr, _ = parse_folder(folder)
    q = ident["cite"][:200] if ident["cite"] else f"{a1} {yr}"
    links = []
    if doi:
        links.append(f"https://doi.org/{doi}")
    links.append("https://www.researchgate.net/search/publication?q=" + urllib.parse.quote(q))
    links.append("https://scholar.google.com/scholar?q=" + urllib.parse.quote(q))
    return " | ".join(links)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--apply", action="store_true", help="actually copy/download PDFs (default: dry run)")
    ap.add_argument("--folder", action="append", help="restrict to these folder(s)")
    ap.add_argument("--scan-dir", action="append", default=[], help="local directory to scan for PDFs")
    ap.add_argument("--no-online", action="store_true", help="skip Crossref / OA lookups")
    ap.add_argument("--no-endnote", action="store_true")
    ap.add_argument("--report", default=os.path.join(CHECKS, "paper_pdf_sources.csv"))
    a = ap.parse_args()

    todo = load_audit(set(a.folder) if a.folder else None)
    print(f"{len(todo)} folder(s) without a paper PDF", file=sys.stderr)
    overrides, registry, sensory = load_overrides(), load_registry(), load_sensory()
    con = None if a.no_endnote else en.connect(en.SDB)

    report = []
    scan_targets = []
    for r in todo:
        f = r["folder"]
        a1, a2, yr, suf = parse_folder(f)
        ident = identify(f, overrides, registry, sensory)
        row = dict(folder=f, status="", source_kind="", source="", endnote_id="", doi=ident["doi"],
                   confidence="", id_source=ident["id_source"], citation=ident["cite"] or ident["title"],
                   endnote_note="", online_note="", manual_links="")

        # EndNote
        if con is not None:
            m = endnote_match(f, ident, con)
            if m:
                row["endnote_note"] = f"[{m['en_id']}] {m['en_cite']} — {m['en_title']} (conf={m['confidence']}, ov={m['overlap']})"
                if m["alternates"]:
                    row["endnote_note"] += f" | alt: {m['alternates']}"
                if not row["doi"] and m["en_doi"] and m["confidence"] in ("doi", "strong"):
                    row["doi"] = doi_in(m["en_doi"])
                if m["confidence"] in ("doi", "strong") and m["pdf"] and os.path.exists(m["pdf"]):
                    status, dst = place_pdf(f, m["pdf"], "endnote", m["pdf"], a.apply,
                                            endnote_id=m["en_id"], doi=row["doi"], confidence=m["confidence"])
                    row.update(status=status, source_kind="endnote", source=m["pdf"], endnote_id=m["en_id"],
                               confidence=m["confidence"])
                elif m["confidence"] in ("doi", "strong"):
                    row["endnote_note"] += " | in EndNote but NO PDF attached"
                    row["endnote_id"] = m["en_id"]
            else:
                row["endnote_note"] = "no EndNote record (first author + year)"

        # Online
        if not row["status"] and not a.no_online:
            if not row["doi"]:
                row["doi"] = crossref_doi(ident["cite"], a1, yr)
                time.sleep(1.0)
            if row["doi"]:
                url, src = oa_pdf(row["doi"])
                if url:
                    try:
                        b, ctype = http_bytes(url)
                        if is_pdf(b):
                            status, dst = place_pdf(f, b, "open_access", url, a.apply, doi=row["doi"], confidence=src)
                            row.update(status=status, source_kind="open_access", source=url, confidence=src)
                        else:
                            row["online_note"] = f"{src}: {url} returned {ctype[:30]} (not a PDF; likely bot check — download by hand)"
                    except Exception as e:
                        row["online_note"] = f"{src}: {url} -> {type(e).__name__}"
                else:
                    row["online_note"] = src
            else:
                row["online_note"] = "no DOI resolved (Crossref)"

        if not row["status"]:
            tt = toks(ident["cite"] or ident["title"])
            scan_targets.append((f, a1, yr, tt))
        row["manual_links"] = manual_links(f, ident, row["doi"])
        report.append(row)

    # Local disk scan for whatever is left
    unchecked = []
    if a.scan_dir and scan_targets:
        hits, unchecked = scan_local(a.scan_dir, scan_targets,
                                     exclude_substrings=[os.sep + "References.Data" + os.sep, ROOT + os.sep])
        for row in report:
            if row["status"] or row["folder"] not in hits:
                continue
            best = sorted(hits[row["folder"]], reverse=True)[0]
            status, dst = place_pdf(row["folder"], best[1], "local_disk", best[1], a.apply, doi=row["doi"],
                                    confidence=f"title_overlap={best[0]}")
            row.update(status=status, source_kind="local_disk", source=best[1], confidence=f"title_overlap={best[0]}")

    for row in report:
        if not row["status"]:
            row["status"] = "NOT FOUND — obtain manually"

    # Write report
    fields = ["folder", "status", "source_kind", "source", "endnote_id", "doi", "confidence", "id_source",
              "citation", "endnote_note", "online_note", "manual_links"]
    with open(a.report, "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=fields)
        w.writeheader()
        w.writerows(report)
    md = a.report[:-4] + ".md"
    with open(md, "w", encoding="utf-8") as fh:
        fh.write(f"# Paper-PDF sources — {time.strftime('%Y-%m-%d')} ({'applied' if a.apply else 'DRY RUN'})\n\n")
        for kind in ("endnote", "open_access", "local_disk", ""):
            sel = [r for r in report if r["source_kind"] == kind]
            if not sel:
                continue
            fh.write(f"## {kind or 'not found'} ({len(sel)})\n\n")
            for r in sel:
                fh.write(f"- **{r['folder']}** — {r['status']}\n")
                if r["source"]:
                    fh.write(f"  - source: `{r['source']}`\n")
                if r["endnote_note"]:
                    fh.write(f"  - EndNote: {r['endnote_note']}\n")
                if r["online_note"]:
                    fh.write(f"  - online: {r['online_note']}\n")
                if r["citation"]:
                    fh.write(f"  - citation: {r['citation'][:220]}\n")
                if not kind:
                    fh.write(f"  - try: {r['manual_links']}\n")
        if unchecked:
            fh.write(f"\n## Local PDFs that could not be read ({len(unchecked)}; cloud placeholders?)\n\n")
            for p in unchecked:
                fh.write(f"- `{p}`\n")
    n = {k: sum(1 for r in report if r["source_kind"] == k) for k in ("endnote", "open_access", "local_disk")}
    n["not_found"] = sum(1 for r in report if not r["source_kind"])
    print(json.dumps(n), file=sys.stderr)
    print(f"report: {a.report}\n        {md}", file=sys.stderr)


if __name__ == "__main__":
    main()
